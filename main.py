import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import time
import threading
import numpy as np
from PIL import Image

from ocr_engine import OCREngine
from question_matcher import QuestionMatcher
from screen_capture import ScreenCapture
from config_manager import ConfigManager
from history_manager import HistoryManager
from clicker import click_at


class MagicQAApp:
    def __init__(self, root):
        self.root = root
        self.root.title("自动化识别显示程序 V3")
        self.root.geometry("640x580")
        self.root.minsize(560, 480)

        self.config = ConfigManager()
        self.matcher = QuestionMatcher()
        self.capture = ScreenCapture()
        self.history = HistoryManager()
        self.ocr = None
        self.ocr_ready = False

        self.auto_running = False
        self.auto_thread = None
        self.last_question_hash = None
        self.last_capture_img = None
        self.is_recognizing = False
        self.last_unmatched_question = None
        self.last_question_region = None

        self._build_ui()
        self._setup_hotkeys()

        if self.config.get('always_on_top', True):
            self.root.attributes('-topmost', True)

        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        self.status_var.set("正在加载OCR引擎...")
        self.root.after(100, self._init_ocr_async)

    def _init_ocr_async(self):
        def init():
            try:
                self.ocr = OCREngine(scale=self.config.get('ocr_scale', 1.5))
                self.ocr_ready = True
                self.root.after(0, lambda: self.status_var.set("就绪 | F8识别 | F9 OCR识别正确答案补充题库"))
            except Exception as e:
                self.root.after(0, lambda: self.status_var.set(f"OCR加载失败: {e}"))
        threading.Thread(target=init, daemon=True).start()

    def _build_ui(self):
        style = ttk.Style()
        try:
            style.theme_use('clam')
        except Exception:
            pass

        main_frame = ttk.Frame(self.root, padding=8)
        main_frame.pack(fill=tk.BOTH, expand=True)

        header = ttk.Frame(main_frame)
        header.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(header, text="自动化识别显示程序", font=("微软雅黑", 15, "bold")).pack(side=tk.LEFT)
        ttk.Label(header, text=f"题库:{len(self.matcher.questions)}题",
                  font=("微软雅黑", 9), foreground="gray").pack(side=tk.RIGHT, padx=5)
        self.auto_status_label = ttk.Label(header, text="自动:关", font=("微软雅黑", 9), foreground="gray")
        self.auto_status_label.pack(side=tk.RIGHT, padx=5)

        notebook = ttk.Notebook(main_frame)
        notebook.pack(fill=tk.BOTH, expand=True)

        self._build_answer_tab(notebook)
        self._build_region_tab(notebook)
        self._build_history_tab(notebook)
        self._build_unmatched_tab(notebook)
        self._build_settings_tab(notebook)
        self._build_question_bank_tab(notebook)

        self.status_var = tk.StringVar(value="就绪 | F8识别 | F9 OCR识别正确答案补充题库")
        status_bar = ttk.Frame(main_frame)
        status_bar.pack(fill=tk.X, pady=(4, 0))
        ttk.Label(status_bar, textvariable=self.status_var, font=("微软雅黑", 9),
                  foreground="gray").pack(side=tk.LEFT)

    def _build_answer_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="答题")

        btn_frame = ttk.Frame(tab)
        btn_frame.pack(fill=tk.X, pady=(0, 8))
        self.recognize_btn = ttk.Button(btn_frame, text="识别答题 (F8)", command=self._recognize)
        self.recognize_btn.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=6, padx=(0, 4))
        self.auto_btn = ttk.Button(btn_frame, text="开启自动识别", command=self._toggle_auto)
        self.auto_btn.pack(side=tk.LEFT, ipady=6, padx=4)
        ttk.Button(btn_frame, text="OCR补充答案 (F9)", command=self._supplement_answer).pack(side=tk.LEFT, ipady=6, padx=4)
        ttk.Button(btn_frame, text="清空", command=self._clear_answer).pack(side=tk.LEFT, ipady=6)

        info_frame = ttk.Frame(tab)
        info_frame.pack(fill=tk.X, pady=(0, 4))
        self.q_type_var = tk.StringVar(value="题目类型: -")
        ttk.Label(info_frame, textvariable=self.q_type_var, font=("微软雅黑", 9),
                  foreground="gray").pack(side=tk.LEFT)
        self.match_info_var = tk.StringVar(value="")
        ttk.Label(info_frame, textvariable=self.match_info_var, font=("微软雅黑", 9),
                  foreground="gray").pack(side=tk.RIGHT)

        answer_frame = ttk.LabelFrame(tab, text="答案", padding=8)
        answer_frame.pack(fill=tk.BOTH, expand=True)
        self.answer_text = tk.Text(answer_frame, font=("微软雅黑", self.config.get('font_size', 18), "bold"),
                                    wrap=tk.WORD, height=5, bg="#FFF8DC", relief=tk.FLAT)
        self.answer_text.pack(fill=tk.BOTH, expand=True)
        self.answer_text.insert(tk.END, "按 F8 识别答题\n按 F9 OCR识别正确答案区域，自动补充到题库")
        self.answer_text.config(state=tk.DISABLED)

        detail_frame = ttk.LabelFrame(tab, text="识别详情", padding=6)
        detail_frame.pack(fill=tk.X, pady=(8, 0))
        self.detail_text = tk.Text(detail_frame, font=("微软雅黑", 9), wrap=tk.WORD, height=3,
                                    bg="#F5F5F5", relief=tk.FLAT)
        self.detail_text.pack(fill=tk.X)
        self.detail_text.config(state=tk.DISABLED)

    def _build_region_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="区域设置")

        ttk.Label(tab, text="题目识别区域（框选题目文字所在位置，建议框大一些）", font=("微软雅黑", 10, "bold")).pack(anchor=tk.W, pady=(0, 4))
        self._build_region_row(tab, 'question_region', '题目')

        ttk.Separator(tab, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=10)

        ttk.Label(tab, text="正确答案区域（F9补充答案时OCR识别此区域）", font=("微软雅黑", 10, "bold")).pack(anchor=tk.W, pady=(0, 4))
        self._build_region_row(tab, 'answer_region', '答案')

        ttk.Separator(tab, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=10)

        ttk.Label(tab, text="选项区域（框选包含A/B/C/D四个选项的区域，开启自动点击后识别此区域匹配答案）", font=("微软雅黑", 10, "bold")).pack(anchor=tk.W, pady=(0, 4))
        self._build_region_row(tab, 'options_region', '选项')

        ttk.Separator(tab, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=10)

        auto_frame = ttk.Frame(tab)
        auto_frame.pack(fill=tk.X)
        ttk.Label(auto_frame, text="自动识别开启后，检测到题目区域画面变化即自动识别。",
                  font=("微软雅黑", 8), foreground="gray").pack(anchor=tk.W, pady=(2, 0))

    def _build_region_row(self, parent, region_key, label):
        frame = ttk.Frame(parent)
        frame.pack(fill=tk.X, pady=2)

        region = self.config.get(region_key, {})
        vars_dict = {}
        for i, (key, lbl) in enumerate([('x', 'X'), ('y', 'Y'), ('w', '宽'), ('h', '高')]):
            ttk.Label(frame, text=lbl + ":").grid(row=0, column=i*2, padx=(4, 1))
            var = tk.StringVar(value=str(region.get(key, 0)))
            vars_dict[key] = var
            ttk.Entry(frame, textvariable=var, width=6).grid(row=0, column=i*2+1, padx=(0, 4))

        setattr(self, f'{region_key}_vars', vars_dict)

        btn_frame = ttk.Frame(parent)
        btn_frame.pack(fill=tk.X, pady=(2, 6))
        ttk.Button(btn_frame, text=f"框选{label}区域",
                   command=lambda k=region_key: self._select_region(k)).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="保存",
                   command=lambda k=region_key: self._save_region(k)).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="测试截图",
                   command=lambda k=region_key: self._test_capture(k)).pack(side=tk.LEFT, padx=2)

    def _build_history_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="识别历史")

        btn_frame = ttk.Frame(tab)
        btn_frame.pack(fill=tk.X, pady=(0, 6))
        ttk.Button(btn_frame, text="刷新", command=self._refresh_history).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="清空历史", command=self._clear_history).pack(side=tk.LEFT, padx=2)

        columns = ('time', 'question', 'answer', 'type', 'conf')
        self.history_tree = ttk.Treeview(tab, columns=columns, show='headings', height=15)
        self.history_tree.heading('time', text='时间')
        self.history_tree.heading('question', text='题目')
        self.history_tree.heading('answer', text='答案')
        self.history_tree.heading('type', text='匹配')
        self.history_tree.heading('conf', text='置信度')
        self.history_tree.column('time', width=110)
        self.history_tree.column('question', width=220)
        self.history_tree.column('answer', width=140)
        self.history_tree.column('type', width=50)
        self.history_tree.column('conf', width=50)
        self.history_tree.pack(fill=tk.BOTH, expand=True)
        self._refresh_history()

    def _build_unmatched_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="未匹配题目")

        btn_frame = ttk.Frame(tab)
        btn_frame.pack(fill=tk.X, pady=(0, 6))
        ttk.Button(btn_frame, text="刷新", command=self._refresh_unmatched).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="导出列表", command=self._export_unmatched).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="清空", command=self._clear_unmatched).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="选中项补充答案", command=self._supplement_selected).pack(side=tk.LEFT, padx=2)

        columns = ('count', 'question', 'first_time', 'last_time')
        self.unmatched_tree = ttk.Treeview(tab, columns=columns, show='headings', height=15)
        self.unmatched_tree.heading('count', text='次数')
        self.unmatched_tree.heading('question', text='题目')
        self.unmatched_tree.heading('first_time', text='首次')
        self.unmatched_tree.heading('last_time', text='最近')
        self.unmatched_tree.column('count', width=50)
        self.unmatched_tree.column('question', width=340)
        self.unmatched_tree.column('first_time', width=110)
        self.unmatched_tree.column('last_time', width=110)
        self.unmatched_tree.pack(fill=tk.BOTH, expand=True)
        self._refresh_unmatched()

    def _build_settings_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="设置")

        frame = ttk.Frame(tab)
        frame.pack(fill=tk.X, pady=4)

        ttk.Label(frame, text="OCR缩放倍数:").grid(row=0, column=0, sticky=tk.W, pady=4)
        self.ocr_scale_var = tk.StringVar(value=str(self.config.get('ocr_scale', 1.5)))
        ttk.Combobox(frame, textvariable=self.ocr_scale_var, values=['1.5', '2.0', '2.5', '3.0'],
                      width=8).grid(row=0, column=1, sticky=tk.W, pady=4)

        ttk.Label(frame, text="匹配阈值:").grid(row=1, column=0, sticky=tk.W, pady=4)
        self.threshold_var = tk.StringVar(value=str(self.config.get('match_threshold', 70)))
        ttk.Combobox(frame, textvariable=self.threshold_var, values=['60', '65', '70', '75', '80'],
                      width=8).grid(row=1, column=1, sticky=tk.W, pady=4)

        ttk.Label(frame, text="答案字号:").grid(row=2, column=0, sticky=tk.W, pady=4)
        self.font_size_var = tk.StringVar(value=str(self.config.get('font_size', 18)))
        ttk.Combobox(frame, textvariable=self.font_size_var, values=['14', '16', '18', '20', '24', '28'],
                      width=8).grid(row=2, column=1, sticky=tk.W, pady=4)

        ttk.Label(frame, text="自动识别间隔(秒):").grid(row=3, column=0, sticky=tk.W, pady=4)
        self.auto_interval_var = tk.StringVar(value=str(self.config.get('auto_interval', 0.3)))
        ttk.Entry(frame, textvariable=self.auto_interval_var, width=8).grid(row=3, column=1, sticky=tk.W, pady=4)

        self.record_history_var = tk.BooleanVar(value=self.config.get('record_history', True))
        ttk.Checkbutton(frame, text="记录识别历史", variable=self.record_history_var).grid(
            row=4, column=0, columnspan=3, sticky=tk.W, pady=4)

        self.record_unmatched_var = tk.BooleanVar(value=self.config.get('record_unmatched', True))
        ttk.Checkbutton(frame, text="记录未匹配题目", variable=self.record_unmatched_var).grid(
            row=5, column=0, columnspan=3, sticky=tk.W, pady=4)

        self.auto_click_var = tk.BooleanVar(value=self.config.get('auto_click', False))
        ttk.Checkbutton(frame, text="自动点击选项（识别到答案后自动点击对应选项，需配置选项区域）", variable=self.auto_click_var).grid(
            row=6, column=0, columnspan=3, sticky=tk.W, pady=4)

        self.always_top_var = tk.BooleanVar(value=self.config.get('always_on_top', True))
        ttk.Checkbutton(frame, text="窗口置顶", variable=self.always_top_var,
                         command=self._toggle_topmost).grid(row=7, column=0, columnspan=3, sticky=tk.W, pady=4)

        ttk.Button(frame, text="保存设置", command=self._save_settings).grid(row=8, column=0, pady=10)

        ttk.Label(frame, text="\n快捷键:\n  F8 = 识别答题\n  F9 = OCR识别正确答案区域，补充到题库\n\n自动识别开启后，检测到题目变化会自动识别。",
                  font=("微软雅黑", 9), foreground="gray").grid(row=9, column=0, columnspan=3, sticky=tk.W, pady=10)

    def _build_question_bank_tab(self, notebook):
        tab = ttk.Frame(notebook, padding=8)
        notebook.add(tab, text="题库管理")

        top_frame = ttk.Frame(tab)
        top_frame.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(top_frame, text="搜索:").pack(side=tk.LEFT)
        self.qb_search_var = tk.StringVar()
        ttk.Entry(top_frame, textvariable=self.qb_search_var, width=30).pack(side=tk.LEFT, padx=5)
        ttk.Button(top_frame, text="搜索", command=self._qb_search).pack(side=tk.LEFT, padx=2)
        ttk.Button(top_frame, text="刷新", command=self._qb_refresh).pack(side=tk.LEFT, padx=2)
        ttk.Button(top_frame, text="导入题库", command=self._qb_import).pack(side=tk.LEFT, padx=2)
        ttk.Button(top_frame, text="导出题库", command=self._qb_export).pack(side=tk.LEFT, padx=2)
        self.qb_count_label = ttk.Label(top_frame, text="共0题", font=("微软雅黑", 9), foreground="gray")
        self.qb_count_label.pack(side=tk.RIGHT)

        list_frame = ttk.Frame(tab)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 6))

        columns = ("question", "answer")
        self.qb_tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=12)
        self.qb_tree.heading("question", text="题目")
        self.qb_tree.heading("answer", text="答案")
        self.qb_tree.column("question", width=380, anchor=tk.W)
        self.qb_tree.column("answer", width=180, anchor=tk.W)
        self.qb_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        qb_scroll = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.qb_tree.yview)
        qb_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.qb_tree.configure(yscrollcommand=qb_scroll.set)
        self.qb_tree.bind("<<TreeviewSelect>>", self._qb_on_select)

        edit_frame = ttk.LabelFrame(tab, text="添加 / 编辑", padding=8)
        edit_frame.pack(fill=tk.X)

        ttk.Label(edit_frame, text="题目:").grid(row=0, column=0, sticky=tk.NW, pady=3)
        self.qb_question_var = tk.StringVar()
        ttk.Entry(edit_frame, textvariable=self.qb_question_var, width=60).grid(row=0, column=1, sticky=tk.W, pady=3)

        ttk.Label(edit_frame, text="答案:").grid(row=1, column=0, sticky=tk.NW, pady=3)
        self.qb_answer_var = tk.StringVar()
        ttk.Entry(edit_frame, textvariable=self.qb_answer_var, width=60).grid(row=1, column=1, sticky=tk.W, pady=3)

        btn_frame = ttk.Frame(edit_frame)
        btn_frame.grid(row=2, column=0, columnspan=2, pady=6)
        ttk.Button(btn_frame, text="添加/保存", command=self._qb_save).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="删除选中", command=self._qb_delete).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="清空输入", command=self._qb_clear).pack(side=tk.LEFT, padx=4)

        self._qb_refresh()

    def _qb_refresh(self):
        for item in self.qb_tree.get_children():
            self.qb_tree.delete(item)
        questions = self.matcher.questions
        for key, item in questions.items():
            q = item.get('question', '')
            a = item.get('answer', '')
            self.qb_tree.insert("", tk.END, values=(q, a))
        self.qb_count_label.config(text=f"共{len(questions)}题")

    def _qb_search(self):
        keyword = self.qb_search_var.get().strip()
        for item in self.qb_tree.get_children():
            self.qb_tree.delete(item)
        questions = self.matcher.questions
        count = 0
        for key, item in questions.items():
            q = item.get('question', '')
            a = item.get('answer', '')
            if not keyword or keyword in q or keyword in a:
                self.qb_tree.insert("", tk.END, values=(q, a))
                count += 1
        self.qb_count_label.config(text=f"匹配{count}题")

    def _qb_on_select(self, event):
        selected = self.qb_tree.selection()
        if not selected:
            return
        values = self.qb_tree.item(selected[0], "values")
        self.qb_question_var.set(values[0])
        self.qb_answer_var.set(values[1])

    def _qb_save(self):
        question = self.qb_question_var.get().strip()
        answer = self.qb_answer_var.get().strip()
        if not question or not answer:
            messagebox.showwarning("提示", "题目和答案都不能为空")
            return
        self._add_to_question_bank(question, answer)
        self._qb_refresh()
        self._qb_clear()
        self.status_var.set("题库已更新")

    def _qb_delete(self):
        selected = self.qb_tree.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选中要删除的题目")
            return
        values = self.qb_tree.item(selected[0], "values")
        question = values[0]
        if not messagebox.askyesno("确认", f"确定删除此题目？\n\n{question}"):
            return
        from question_matcher import normalize_question
        key = normalize_question(question)
        if key in self.matcher.questions:
            del self.matcher.questions[key]
            self._save_question_bank()
            self._qb_refresh()
            self._qb_clear()
            self.status_var.set("题目已删除")

    def _qb_clear(self):
        self.qb_question_var.set("")
        self.qb_answer_var.set("")
        self.qb_tree.selection_remove(self.qb_tree.selection())

    def _qb_import(self):
        from tkinter import filedialog
        file_path = filedialog.askopenfilename(
            title="选择题库文件",
            filetypes=[("题库文件", "*.json *.xlsx *.xls *.csv"), ("JSON文件", "*.json"), ("Excel文件", "*.xlsx *.xls"), ("CSV文件", "*.csv"), ("所有文件", "*.*")]
        )
        if not file_path:
            return
        try:
            added = 0
            updated = 0
            if file_path.lower().endswith('.json'):
                added, updated = self._import_json(file_path)
            elif file_path.lower().endswith(('.xlsx', '.xls')):
                added, updated = self._import_excel(file_path)
            elif file_path.lower().endswith('.csv'):
                added, updated = self._import_csv(file_path)
            else:
                messagebox.showerror("错误", "不支持的文件格式")
                return
            self._save_question_bank()
            self._qb_refresh()
            messagebox.showinfo("导入完成", f"导入成功！\n\n新增: {added} 题\n更新: {updated} 题\n当前题库: {len(self.matcher.questions)} 题")
            self.status_var.set(f"导入完成：新增{added}题，更新{updated}题")
        except Exception as e:
            messagebox.showerror("导入失败", str(e))
            import traceback
            traceback.print_exc()

    def _import_json(self, file_path):
        import json
        from question_matcher import normalize_question, clean_text
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        questions = data.get('questions', data)
        added = 0
        updated = 0
        for key, item in questions.items():
            if isinstance(item, dict):
                q = item.get('question', '')
                a = item.get('answer', '')
            else:
                q = key
                a = str(item)
            if q and a:
                norm_key = normalize_question(q)
                if norm_key in self.matcher.questions:
                    updated += 1
                else:
                    added += 1
                self.matcher.questions[norm_key] = {'question': clean_text(q), 'answer': clean_text(a)}
        return added, updated

    def _import_excel(self, file_path):
        from question_matcher import normalize_question, clean_text
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True)
        added = 0
        updated = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c).strip().lower() if c else '' for c in rows[0]]
            q_col = 0
            a_col = 1
            for i, h in enumerate(header):
                if '题' in h or 'question' in h:
                    q_col = i
                if '答' in h or 'answer' in h:
                    a_col = i
            for row in rows[1:]:
                if not row or len(row) <= max(q_col, a_col):
                    continue
                q = str(row[q_col]).strip() if row[q_col] else ''
                a = str(row[a_col]).strip() if row[a_col] else ''
                if q and a and q != 'None' and a != 'None':
                    norm_key = normalize_question(q)
                    if norm_key in self.matcher.questions:
                        updated += 1
                    else:
                        added += 1
                    self.matcher.questions[norm_key] = {'question': clean_text(q), 'answer': clean_text(a)}
        wb.close()
        return added, updated

    def _import_csv(self, file_path):
        import csv
        from question_matcher import normalize_question, clean_text
        added = 0
        updated = 0
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            rows = list(reader)
        if not rows:
            return 0, 0
        header = [c.strip().lower() for c in rows[0]]
        q_col = 0
        a_col = 1
        for i, h in enumerate(header):
            if '题' in h or 'question' in h:
                q_col = i
            if '答' in h or 'answer' in h:
                a_col = i
        for row in rows[1:]:
            if not row or len(row) <= max(q_col, a_col):
                continue
            q = row[q_col].strip()
            a = row[a_col].strip()
            if q and a:
                norm_key = normalize_question(q)
                if norm_key in self.matcher.questions:
                    updated += 1
                else:
                    added += 1
                self.matcher.questions[norm_key] = {'question': clean_text(q), 'answer': clean_text(a)}
        return added, updated

    def _qb_export(self):
        from tkinter import filedialog
        file_path = filedialog.asksaveasfilename(
            title="导出题库",
            defaultextension=".json",
            filetypes=[("JSON文件", "*.json"), ("Excel文件", "*.xlsx"), ("CSV文件", "*.csv")]
        )
        if not file_path:
            return
        try:
            if file_path.lower().endswith('.json'):
                import json
                data = {'version': '5.0', 'total': len(self.matcher.questions), 'questions': self.matcher.questions}
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
            elif file_path.lower().endswith('.xlsx'):
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "题库"
                ws.append(["题目", "答案"])
                for key, item in self.matcher.questions.items():
                    ws.append([item.get('question', ''), item.get('answer', '')])
                wb.save(file_path)
            elif file_path.lower().endswith('.csv'):
                import csv
                with open(file_path, 'w', encoding='utf-8-sig', newline='') as f:
                    writer = csv.writer(f)
                    writer.writerow(["题目", "答案"])
                    for key, item in self.matcher.questions.items():
                        writer.writerow([item.get('question', ''), item.get('answer', '')])
            messagebox.showinfo("导出完成", f"题库已导出到:\n{file_path}")
            self.status_var.set("题库已导出")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _save_question_bank(self):
        import json
        data = {
            "version": "5.0",
            "total": len(self.matcher.questions),
            "questions": self.matcher.questions
        }
        path = self.matcher.questions_path
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def _setup_hotkeys(self):
        try:
            self.root.bind('<F8>', lambda e: self._recognize())
            self.root.bind('<F9>', lambda e: self._supplement_answer())
        except Exception:
            pass

    def _toggle_topmost(self):
        self.root.attributes('-topmost', self.always_top_var.get())
        self.config.set('always_on_top', self.always_top_var.get())

    def _save_settings(self):
        try:
            self.config.set('ocr_scale', float(self.ocr_scale_var.get()))
            self.config.set('match_threshold', int(self.threshold_var.get()))
            self.config.set('font_size', int(self.font_size_var.get()))
            self.config.set('auto_interval', float(self.auto_interval_var.get()))
            self.config.set('record_history', self.record_history_var.get())
            self.config.set('record_unmatched', self.record_unmatched_var.get())
            self.config.set('auto_click', self.auto_click_var.get())
            self.ocr.scale = float(self.ocr_scale_var.get())
            self.answer_text.config(font=("微软雅黑", int(self.font_size_var.get()), "bold"))
            self.status_var.set("设置已保存")
            messagebox.showinfo("提示", "设置已保存")
        except ValueError:
            messagebox.showerror("错误", "请输入有效的数值")

    def _select_region(self, region_key):
        result = {'x': 0, 'y': 0, 'w': 0, 'h': 0, 'ok': False}
        try:
            sel_root = tk.Toplevel(self.root)
            sel_root.attributes('-fullscreen', True)
            sel_root.attributes('-alpha', 0.3)
            sel_root.attributes('-topmost', True)
            sel_root.configure(bg='black')
            sel_root.overrideredirect(True)
            canvas = tk.Canvas(sel_root, cursor='cross', bg='black', highlightthickness=0)
            canvas.pack(fill=tk.BOTH, expand=True)
            start_x = start_y = 0
            rect = None

            def on_press(e):
                nonlocal start_x, start_y, rect
                start_x, start_y = e.x, e.y
                rect = canvas.create_rectangle(start_x, start_y, start_x, start_y, outline='red', width=2)

            def on_drag(e):
                if rect:
                    canvas.coords(rect, start_x, start_y, e.x, e.y)

            def on_release(e):
                x1, y1 = min(start_x, e.x), min(start_y, e.y)
                x2, y2 = max(start_x, e.x), max(start_y, e.y)
                w, h = x2 - x1, y2 - y1
                if w > 10 and h > 10:
                    result['x'] = x1
                    result['y'] = y1
                    result['w'] = w
                    result['h'] = h
                    result['ok'] = True
                sel_root.destroy()

            canvas.bind('<ButtonPress-1>', on_press)
            canvas.bind('<B1-Motion>', on_drag)
            canvas.bind('<ButtonRelease-1>', on_release)
            sel_root.bind('<Escape>', lambda e: sel_root.destroy())
            sel_root.wait_window()
        except Exception as e:
            messagebox.showerror("错误", f"框选失败: {e}")

        if result['ok']:
            vars_dict = getattr(self, f'{region_key}_vars', {})
            vars_dict['x'].set(str(result['x']))
            vars_dict['y'].set(str(result['y']))
            vars_dict['w'].set(str(result['w']))
            vars_dict['h'].set(str(result['h']))
            self._save_region(region_key)

        self.root.lift()
        self.root.focus_force()

    def _save_region(self, region_key):
        try:
            vars_dict = getattr(self, f'{region_key}_vars', {})
            x = int(vars_dict['x'].get())
            y = int(vars_dict['y'].get())
            w = int(vars_dict['w'].get())
            h = int(vars_dict['h'].get())
            self.config.set_region(region_key, x, y, w, h)
            self.status_var.set(f"区域已保存: ({x},{y}) {w}x{h}")
        except ValueError:
            messagebox.showerror("错误", "坐标必须为整数")

    def _test_capture(self, region_key):
        try:
            region = self.config.get_region(region_key)
            if not region:
                messagebox.showerror("错误", "请先设置区域")
                return
            x, y, w, h = region
            img = self.capture.capture_region(x, y, w, h)
            import os
            save_dir = os.path.dirname(os.path.abspath(__file__))
            path = os.path.join(save_dir, f'test_{region_key}.png')
            img.save(path)
            self._show_image_preview(img, f"截图预览 - {region_key} ({img.width}x{img.height})")
            self.status_var.set(f"截图已保存: {path}")
        except Exception as e:
            messagebox.showerror("错误", f"截图失败: {e}")

    def _show_image_preview(self, pil_img, title):
        try:
            from PIL import ImageTk
            top = tk.Toplevel(self.root)
            top.title(title)
            top.resizable(False, False)
            max_w = 800
            max_h = 600
            img_w, img_h = pil_img.size
            scale = min(max_w / img_w, max_h / img_h, 1.0)
            if scale < 1.0:
                pil_img = pil_img.resize((int(img_w * scale), int(img_h * scale)), Image.LANCZOS)
            photo = ImageTk.PhotoImage(pil_img)
            label = tk.Label(top, image=photo)
            label.image = photo
            label.pack()
            top.transient(self.root)
            top.grab_set()
            top.focus_set()
            top.bind('<Escape>', lambda e: top.destroy())
            top.after(3000, lambda: top.destroy() if top.winfo_exists() else None)
        except Exception as e:
            self.status_var.set(f"预览显示失败: {e}")

    def _toggle_auto(self):
        if self.auto_running:
            self.auto_running = False
            self.auto_btn.config(text="开启自动识别")
            self.auto_status_label.config(text="自动:关")
            self.status_var.set("自动识别已关闭")
        else:
            self.auto_running = True
            self.auto_btn.config(text="停止自动识别")
            self.auto_status_label.config(text="自动:开")
            self.status_var.set("自动识别已开启，等待题目出现...")
            self.auto_thread = threading.Thread(target=self._auto_loop, daemon=True)
            self.auto_thread.start()

    def _auto_loop(self):
        recognize_start = 0
        while self.auto_running:
            try:
                interval = self.config.get('auto_interval', 0.3)
                if self.is_recognizing and time.time() - recognize_start > 10:
                    self.is_recognizing = False
                    self.root.after(0, lambda: self.recognize_btn.config(state=tk.NORMAL))
                    self.root.after(0, lambda: self.status_var.set("识别超时，已重置"))
                region = self.config.get_region('question_region')
                if not region:
                    time.sleep(interval)
                    continue
                x, y, w, h = region
                img = self.capture.capture_region(x, y, w, h)
                img_hash = self._image_hash(img)
                if img_hash != self.last_question_hash:
                    self.last_question_hash = img_hash
                    if not self.is_recognizing:
                        self.is_recognizing = True
                        recognize_start = time.time()
                        self.root.after(0, lambda: self.recognize_btn.config(state=tk.DISABLED))
                        self.root.after(0, lambda: self.status_var.set("识别中..."))
                        threading.Thread(target=self._do_recognize, daemon=True).start()
                time.sleep(interval)
            except Exception as e:
                print(f"自动识别循环错误: {e}")
                time.sleep(0.5)

    def _image_hash(self, img):
        small = img.resize((8, 8)).convert('L')
        arr = np.array(small)
        return hash(arr.tobytes())

    def _recognize(self):
        if self.is_recognizing:
            return
        if not self.ocr_ready:
            self.status_var.set("OCR引擎加载中，请稍候...")
            return
        self.is_recognizing = True
        self.recognize_btn.config(state=tk.DISABLED)
        self.status_var.set("识别中...")
        threading.Thread(target=self._do_recognize, daemon=True).start()

    def _do_recognize(self):
        try:
            t0 = time.time()
            region = self.config.get_region('question_region')

            if not region:
                self._show_answer("请先在「区域设置」中设置题目识别区域", "错误")
                return

            x, y, w, h = region
            img = self.capture.capture_region(x, y, w, h)
            q_result = self.ocr.recognize_question(img)
            q_text = q_result['full_text']
            line_count = q_result.get('line_count', 1)

            t1 = time.time()
            ocr_time = (t1 - t0) * 1000

            self.root.after(0, lambda: self.q_type_var.set(f"{line_count}行 OCR:{ocr_time:.0f}ms"))

            if not q_text:
                self._show_answer("未识别到题目文字\n\n请检查区域设置或调整OCR缩放倍数", "未识别")
                self._show_detail(f"OCR耗时: {ocr_time:.0f}ms\n未识别到文字")
                self.last_unmatched_question = None
                return

            match = self.matcher.match(q_text, threshold=self.config.get('match_threshold', 70))

            if match:
                answer = match['answer']
                conf = match['confidence']
                match_type = "精确" if match['match_type'] == 'exact' else "模糊"
                self._show_answer(answer, f"{match_type}匹配 {conf:.0%}")
                self._show_detail(f"识别题目: {q_text}\n匹配类型: {match_type}\n置信度: {conf:.2%}\nOCR耗时: {ocr_time:.0f}ms")
                self.root.after(0, lambda: self.match_info_var.set(f"{match_type} {conf:.0%}"))
                self.last_unmatched_question = None
                if self.config.get('record_history', True):
                    self.history.add_record(q_text, answer, match_type, conf, 'local', ocr_time)
                if self.config.get('auto_click', False):
                    self._click_matched_option(answer)
            else:
                self.last_unmatched_question = q_text
                self._show_answer(f"题库未找到答案\n\n识别题目:\n{q_text}\n\n请手动作答，然后按 F9 补充正确答案到题库", "未匹配")
                self._show_detail(f"识别题目: {q_text}\n题库未匹配\n按 F9 可补充正确答案\nOCR耗时: {ocr_time:.0f}ms")
                self.root.after(0, lambda: self.match_info_var.set("未匹配 - 按F9补充"))
                if self.config.get('record_unmatched', True):
                    self.history.add_unmatched(q_text, q_text, 'not_in_bank')
                if self.config.get('record_history', True):
                    self.history.add_record(q_text, '(未匹配)', 'none', 0, 'local', ocr_time)

            self.root.after(0, self._refresh_history)
            self.root.after(0, self._refresh_unmatched)
        except Exception as e:
            self._show_answer(f"识别出错: {e}", "错误")
            import traceback
            traceback.print_exc()
        finally:
            self.is_recognizing = False
            self.root.after(0, lambda: self.recognize_btn.config(state=tk.NORMAL))

    def _click_matched_option(self, answer):
        try:
            options_region = self.config.get_region('options_region')
            if not options_region:
                return
            ox, oy, ow, oh = options_region
            img = self.capture.capture_region(ox, oy, ow, oh)
            options = self.ocr.recognize_options(img)
            if not options:
                return
            from rapidfuzz import fuzz
            best_score = 0
            best_option = None
            for opt in options:
                score = fuzz.ratio(answer, opt['text'])
                if score > best_score:
                    best_score = score
                    best_option = opt
            if best_option and best_score >= 50:
                click_x = ox + best_option['center_x']
                click_y = oy + best_option['center_y']
                click_at(click_x, click_y, dwell_min=20, dwell_max=50, random_offset=8)
                self.root.after(0, lambda: self.status_var.set(f"已点击选项: {best_option['text'][:20]} (匹配度{best_score:.0f}%)"))
        except Exception as e:
            print(f"自动点击出错: {e}")

    def _supplement_answer(self):
        if not self.ocr_ready:
            self.status_var.set("OCR引擎加载中，请稍候...")
            return
        if not self.last_unmatched_question:
            messagebox.showinfo("提示", "没有待补充的未匹配题目\n\n请先识别一道题库中没有的题目")
            return

        question = self.last_unmatched_question
        answer_region = self.config.get_region('answer_region')
        if not answer_region:
            answer = simpledialog.askstring("补充正确答案", f"题目：\n{question}\n\n未设置答案区域，请手动输入正确答案：", parent=self.root)
            if answer and answer.strip():
                self._add_to_question_bank(question, answer.strip())
                self.last_unmatched_question = None
                self._show_answer(f"已补充到题库！\n\n题目: {question}\n答案: {answer.strip()}", "已补充")
                self.status_var.set("答案已补充到题库")
                self._refresh_unmatched()
            return

        threading.Thread(target=self._ocr_and_supplement, args=(question, answer_region), daemon=True).start()

    def _ocr_and_supplement(self, question, answer_region):
        try:
            self.root.after(0, lambda: self.status_var.set("正在OCR识别正确答案..."))
            x, y, w, h = answer_region
            img = self.capture.capture_region(x, y, w, h)
            q_result = self.ocr.recognize_question(img)
            answer = q_result['full_text'].strip()

            if answer:
                self._add_to_question_bank(question, answer)
                self.last_unmatched_question = None
                self._show_answer(f"已补充到题库！\n\n题目: {question}\n答案: {answer}", "已补充(OCR)")
                self.root.after(0, lambda: self.status_var.set("OCR识别成功，答案已补充到题库"))
                self.root.after(0, self._refresh_unmatched)
            else:
                self.root.after(0, lambda: self._manual_supplement_fallback(question))
        except Exception as e:
            self.root.after(0, lambda: self._manual_supplement_fallback(question))

    def _manual_supplement_fallback(self, question):
        answer = simpledialog.askstring("OCR识别失败，请手动输入", f"题目：\n{question}\n\n请输入正确答案：", parent=self.root)
        if answer and answer.strip():
            self._add_to_question_bank(question, answer.strip())
            self.last_unmatched_question = None
            self._show_answer(f"已补充到题库！\n\n题目: {question}\n答案: {answer.strip()}", "已补充(手动)")
            self.status_var.set("答案已补充到题库")
            self._refresh_unmatched()

    def _supplement_selected(self):
        selection = self.unmatched_tree.selection()
        if not selection:
            messagebox.showinfo("提示", "请先在列表中选中一道题目")
            return
        item = self.unmatched_tree.item(selection[0])
        question = item['values'][1]
        answer = simpledialog.askstring("补充正确答案", f"题目：\n{question}\n\n请输入正确答案：", parent=self.root)
        if answer and answer.strip():
            self._add_to_question_bank(question, answer.strip())
            self._refresh_unmatched()
            messagebox.showinfo("成功", f"已补充到题库！\n\n答案: {answer.strip()}")
        elif answer is not None:
            messagebox.showwarning("提示", "答案不能为空")

    def _add_to_question_bank(self, question, answer):
        import re
        questions_path = self.matcher.questions_path
        try:
            with open(questions_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception:
            data = {'version': '5.0', 'total': 0, 'questions': {}}

        def normalize(text):
            text = str(text).strip()
            text = re.sub(r'\s+', '', text)
            text = re.sub(r'[]?？,，.。!！:：;；""''「」《》()（）[【】、/|]', '', text)
            return text.lower()

        def clean(text):
            text = str(text).strip()
            text = re.sub(r'\s+', '', text)
            return text

        q_norm = normalize(question)
        if q_norm:
            data['questions'][q_norm] = {'question': clean(question), 'answer': clean(answer)}
            data['total'] = len(data['questions'])
            with open(questions_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            self.matcher.questions = data['questions']
            return True
        return False

    def _show_answer(self, answer, subtitle=""):
        def update():
            self.answer_text.config(state=tk.NORMAL)
            self.answer_text.delete(1.0, tk.END)
            if subtitle:
                self.answer_text.insert(tk.END, f"【{subtitle}】\n\n")
            self.answer_text.insert(tk.END, answer)
            self.answer_text.config(state=tk.DISABLED)
        self.root.after(0, update)

    def _show_detail(self, text):
        def update():
            self.detail_text.config(state=tk.NORMAL)
            self.detail_text.delete(1.0, tk.END)
            self.detail_text.insert(tk.END, text)
            self.detail_text.config(state=tk.DISABLED)
        self.root.after(0, update)

    def _clear_answer(self):
        self.answer_text.config(state=tk.NORMAL)
        self.answer_text.delete(1.0, tk.END)
        self.answer_text.insert(tk.END, "按 F8 识别答题\n按 F9 OCR识别正确答案区域，自动补充到题库")
        self.answer_text.config(state=tk.DISABLED)
        self.q_type_var.set("题目类型: -")
        self.match_info_var.set("")
        self.last_question_hash = None
        self.last_capture_img = None

    def _refresh_history(self):
        for item in self.history_tree.get_children():
            self.history_tree.delete(item)
        for r in self.history.get_recent(50):
            self.history_tree.insert('', tk.END, values=(
                r.get('time', '')[-8:],
                r.get('question', '')[:30],
                r.get('answer', '')[:20],
                r.get('match_type', ''),
                f"{r.get('confidence', 0):.0%}" if r.get('confidence') else '-'
            ))

    def _clear_history(self):
        if messagebox.askyesno("确认", "确定清空所有识别历史？"):
            self.history.clear_history()
            self._refresh_history()

    def _refresh_unmatched(self):
        for item in self.unmatched_tree.get_children():
            self.unmatched_tree.delete(item)
        for r in self.history.get_unmatched(100):
            q = r.get('question', '')
            if q and q != '(空)':
                self.unmatched_tree.insert('', tk.END, values=(
                    r.get('count', 1),
                    q[:50],
                    r.get('time', '')[-8:],
                    r.get('last_time', r.get('time', ''))[-8:]
                ))

    def _export_unmatched(self):
        path = filedialog.asksaveasfilename(defaultextension=".txt",
                                             filetypes=[("文本文件", "*.txt")],
                                             initialfile="未匹配题目列表.txt")
        if path:
            if self.history.export_unmatched(path):
                messagebox.showinfo("成功", f"已导出到: {path}")
            else:
                messagebox.showerror("失败", "导出失败")

    def _clear_unmatched(self):
        if messagebox.askyesno("确认", "确定清空所有未匹配题目记录？"):
            self.history.clear_unmatched()
            self._refresh_unmatched()

    def _on_close(self):
        self.auto_running = False
        self.root.destroy()


import json


def main():
    root = tk.Tk()
    app = MagicQAApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
